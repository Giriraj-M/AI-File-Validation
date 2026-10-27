import os
import csv
import zipfile
from pathlib import Path
from config import Config
from pypdf import PdfReader
from pypdf.errors import PdfReadError, EmptyFileError, PyPdfError
import docx
import openpyxl

class IntegrityService:
    """
    Dedicated file integrity and corruption verification service.
    Performs deep format-specific structural tests across PDF, DOCX, XLSX, PPTX,
    Images, Text, CSV, and ZIP files to ensure documents are intact and undamaged.
    """

    @classmethod
    def check_integrity(cls, file_path: str, detected_category: str, original_filename: str, allowed_types: list[str] = None) -> dict:
        """
        Validates whether the file is structurally intact or corrupted.
        Resolves the expected target format from detected category, filename extension,
        or user prompt requirements.
        """
        if not os.path.exists(file_path):
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": "File does not exist or was removed before inspection.",
                "details": {"error_type": "FileNotFound"}
            }

        file_size = os.path.getsize(file_path)
        if file_size == 0:
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": "File is completely empty (0 bytes) and contains no data stream.",
                "details": {"error_type": "EmptyFile", "size_bytes": 0}
            }

        target_category = detected_category
        ext = Path(original_filename).suffix.lower()

        # If detected_category is unknown, infer expected category from filename extension or user prompt
        if target_category == "unknown":
            for cat, conf in Config.FILE_TYPES.items():
                if ext in conf.get("extensions", []):
                    target_category = cat
                    break
            if target_category == "unknown" and allowed_types:
                for req_t in allowed_types:
                    req_c = req_t.lower().strip(".")
                    for cat, conf in Config.FILE_TYPES.items():
                        if req_c == cat or f".{req_c}" in conf.get("extensions", []):
                            target_category = cat
                            break
                    if target_category != "unknown":
                        break

        try:
            if target_category == "pdf":
                return cls._check_pdf_integrity(file_path)
            elif target_category == "word":
                return cls._check_word_integrity(file_path, original_filename)
            elif target_category == "excel":
                return cls._check_excel_integrity(file_path, original_filename)
            elif target_category == "powerpoint":
                return cls._check_powerpoint_integrity(file_path, original_filename)
            elif target_category == "image":
                return cls._check_image_integrity(file_path)
            elif target_category == "zip":
                return cls._check_zip_integrity(file_path)
            elif target_category in ["text", "csv"]:
                return cls._check_text_csv_integrity(file_path, target_category)
            else:
                return {
                    "is_corrupted": True,
                    "status": "corrupted",
                    "reason": f"File content does not match any recognized or supported document structure (unidentified binary stream '{ext or 'no-ext'}').",
                    "details": {"detected_category": "unknown", "extension": ext}
                }
        except Exception as e:
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": f"Unexpected error during structural integrity check: {str(e)}",
                "details": {"error": str(e)}
            }

    @classmethod
    def _check_pdf_integrity(cls, file_path: str) -> dict:
        """Inspects PDF header, xref table, catalog, trailer, and page streams."""
        try:
            with open(file_path, "rb") as f:
                header = f.read(1024)
            if not header.startswith(b"%PDF-") and b"%PDF-" not in header[:1024]:
                return {
                    "is_corrupted": True,
                    "status": "corrupted",
                    "reason": "Corrupted PDF file: Missing valid '%PDF-' header signature.",
                    "details": {"error_type": "MissingHeader"}
                }

            reader = PdfReader(file_path)
            
            # Password-protected / encrypted PDF is intact even if locked
            if reader.is_encrypted:
                page_count = len(reader.pages) if hasattr(reader, "pages") else 0
                return {
                    "is_corrupted": False,
                    "status": "intact",
                    "reason": "PDF file structure is intact (encrypted/password-protected).",
                    "details": {"pages": page_count, "is_encrypted": True}
                }

            page_count = len(reader.pages)
            if page_count == 0:
                return {
                    "is_corrupted": True,
                    "status": "corrupted",
                    "reason": "PDF file has no readable pages or contains an empty document catalog.",
                    "details": {"pages": 0}
                }

            # Verify that the first page object can be resolved and accessed
            _ = reader.pages[0].extract_text()

            return {
                "is_corrupted": False,
                "status": "intact",
                "reason": f"PDF structure and xref tables are intact ({page_count} page{'s' if page_count != 1 else ''}).",
                "details": {"pages": page_count, "is_encrypted": False}
            }
        except (PyPdfError, PdfReadError, EmptyFileError) as e:
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": f"Corrupted or invalid PDF structure: {str(e)}",
                "details": {"error": str(e), "error_type": type(e).__name__}
            }
        except Exception as e:
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": f"Corrupted PDF file stream: {str(e)}",
                "details": {"error": str(e)}
            }

    @classmethod
    def _check_word_integrity(cls, file_path: str, filename: str) -> dict:
        """Validates DOCX OpenXML package and internal XML structure or legacy DOC."""
        ext = Path(filename).suffix.lower()
        if ext == ".docx":
            try:
                # 1. Test ZIP container CRC and integrity
                with zipfile.ZipFile(file_path, "r") as z:
                    bad_member = z.testzip()
                    if bad_member:
                        return {
                            "is_corrupted": True,
                            "status": "corrupted",
                            "reason": f"Corrupted ZIP package CRC in Word document (damaged part: {bad_member}).",
                            "details": {"corrupt_part": bad_member}
                        }
                    namelist_lower = [n.lower() for n in z.namelist()]
                    if not any("word/document.xml" in n for n in namelist_lower) and not any("[content_types].xml" in n for n in namelist_lower):
                        return {
                            "is_corrupted": True,
                            "status": "corrupted",
                            "reason": "Missing required OpenXML Word parts in package.",
                            "details": {}
                        }

                # 2. Test python-docx parsing
                doc = docx.Document(file_path)
                p_count = len(doc.paragraphs)
                return {
                    "is_corrupted": False,
                    "status": "intact",
                    "reason": f"Word document structure and XML elements are intact ({p_count} paragraph{'s' if p_count != 1 else ''}).",
                    "details": {"paragraphs": p_count}
                }
            except (zipfile.BadZipFile, Exception) as e:
                return {
                    "is_corrupted": True,
                    "status": "corrupted",
                    "reason": f"Corrupted Word document structure: {str(e)}",
                    "details": {"error": str(e)}
                }
        else:
            # Legacy DOC
            with open(file_path, "rb") as f:
                hdr = f.read(8)
            if hdr.startswith(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"):
                return {
                    "is_corrupted": False,
                    "status": "intact",
                    "reason": "Legacy Word compound document (OLE2) structure is intact.",
                    "details": {}
                }
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": "Corrupted or damaged legacy Word document (invalid OLE header).",
                "details": {}
            }

    @classmethod
    def _check_excel_integrity(cls, file_path: str, filename: str) -> dict:
        """Validates XLSX OpenXML package and internal workbook XML or legacy XLS."""
        ext = Path(filename).suffix.lower()
        if ext == ".xlsx":
            try:
                with zipfile.ZipFile(file_path, "r") as z:
                    bad_member = z.testzip()
                    if bad_member:
                        return {
                            "is_corrupted": True,
                            "status": "corrupted",
                            "reason": f"Corrupted ZIP package CRC in Excel spreadsheet (damaged part: {bad_member}).",
                            "details": {"corrupt_part": bad_member}
                        }

                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                sheets = wb.sheetnames
                wb.close()
                return {
                    "is_corrupted": False,
                    "status": "intact",
                    "reason": f"Excel workbook structure and sheets are intact ({len(sheets)} sheet{'s' if len(sheets) != 1 else ''}).",
                    "details": {"sheet_names": sheets}
                }
            except (zipfile.BadZipFile, Exception) as e:
                return {
                    "is_corrupted": True,
                    "status": "corrupted",
                    "reason": f"Corrupted Excel workbook structure: {str(e)}",
                    "details": {"error": str(e)}
                }
        else:
            # Legacy XLS
            with open(file_path, "rb") as f:
                hdr = f.read(8)
            if hdr.startswith(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"):
                return {
                    "is_corrupted": False,
                    "status": "intact",
                    "reason": "Legacy Excel compound document (OLE2) structure is intact.",
                    "details": {}
                }
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": "Corrupted legacy Excel workbook (invalid OLE header).",
                "details": {}
            }

    @classmethod
    def _check_powerpoint_integrity(cls, file_path: str, filename: str) -> dict:
        """Validates PPTX OpenXML package structure."""
        ext = Path(filename).suffix.lower()
        if ext == ".pptx":
            try:
                with zipfile.ZipFile(file_path, "r") as z:
                    bad_member = z.testzip()
                    if bad_member:
                        return {
                            "is_corrupted": True,
                            "status": "corrupted",
                            "reason": f"Corrupted ZIP package CRC in PowerPoint presentation ({bad_member}).",
                            "details": {}
                        }
                    namelist_lower = [n.lower() for n in z.namelist()]
                    if not any("ppt/presentation.xml" in n for n in namelist_lower) and not any("[content_types].xml" in n for n in namelist_lower):
                        return {
                            "is_corrupted": True,
                            "status": "corrupted",
                            "reason": "Missing required OpenXML PowerPoint parts.",
                            "details": {}
                        }
                return {
                    "is_corrupted": False,
                    "status": "intact",
                    "reason": "PowerPoint presentation archive and OpenXML structure are intact.",
                    "details": {}
                }
            except Exception as e:
                return {
                    "is_corrupted": True,
                    "status": "corrupted",
                    "reason": f"Corrupted PowerPoint presentation: {str(e)}",
                    "details": {"error": str(e)}
                }
        else:
            with open(file_path, "rb") as f:
                hdr = f.read(8)
            if hdr.startswith(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"):
                return {
                    "is_corrupted": False,
                    "status": "intact",
                    "reason": "Legacy PowerPoint compound document structure is intact.",
                    "details": {}
                }
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": "Corrupted legacy PowerPoint presentation (invalid OLE header).",
                "details": {}
            }

    @classmethod
    def _check_image_integrity(cls, file_path: str) -> dict:
        """Tests image header, structural blocks, and pixel stream integrity via Pillow."""
        try:
            from PIL import Image, UnidentifiedImageError

            # Pass 1: Verify header and structure
            with Image.open(file_path) as img:
                img.verify()

            # Pass 2: Load pixel data to catch truncated byte streams (verify only reads headers)
            with Image.open(file_path) as img:
                img.load()
                w, h = img.size
                fmt = img.format or "Image"

            return {
                "is_corrupted": False,
                "status": "intact",
                "reason": f"Image binary structure and pixel streams are intact ({fmt}, {w}x{h} px).",
                "details": {"format": fmt, "width": w, "height": h}
            }
        except (UnidentifiedImageError, OSError, Exception) as e:
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": f"Corrupted or truncated image file: {str(e)}",
                "details": {"error": str(e)}
            }

    @classmethod
    def _check_zip_integrity(cls, file_path: str) -> dict:
        """Verifies ZIP central directory and individual entry CRC checksums."""
        try:
            with zipfile.ZipFile(file_path, "r") as z:
                bad_entry = z.testzip()
                if bad_entry:
                    return {
                        "is_corrupted": True,
                        "status": "corrupted",
                        "reason": f"Damaged entry CRC checksum detected in ZIP archive: '{bad_entry}'.",
                        "details": {"bad_entry": bad_entry}
                    }
                count = len(z.infolist())
                return {
                    "is_corrupted": False,
                    "status": "intact",
                    "reason": f"ZIP archive container and member checksums verified ({count} entries).",
                    "details": {"entry_count": count}
                }
        except (zipfile.BadZipFile, Exception) as e:
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": f"Corrupted or incomplete ZIP archive: {str(e)}",
                "details": {"error": str(e)}
            }

    @classmethod
    def _check_text_csv_integrity(cls, file_path: str, category: str) -> dict:
        """Validates character encoding and row structure for plain text and CSV."""
        try:
            with open(file_path, "rb") as f:
                content = f.read(131072)  # Read up to 128KB sample

            # Check for embedded null bytes (indicates binary corruption or payload injection)
            if b"\x00" in content:
                return {
                    "is_corrupted": True,
                    "status": "corrupted",
                    "reason": "Text file contains illegal binary null bytes (corrupted or disguised binary content).",
                    "details": {"has_null_bytes": True}
                }

            # Verify UTF-8 or Latin-1 decoding
            try:
                decoded = content.decode("utf-8")
            except UnicodeDecodeError:
                try:
                    decoded = content.decode("latin-1")
                except Exception:
                    return {
                        "is_corrupted": True,
                        "status": "corrupted",
                        "reason": "Text file contains invalid or corrupted character encoding.",
                        "details": {}
                    }

            if category == "csv":
                lines = decoded.splitlines()
                if not lines:
                    return {
                        "is_corrupted": True,
                        "status": "corrupted",
                        "reason": "CSV spreadsheet has no readable lines.",
                        "details": {}
                    }
                reader = csv.reader(lines[:100])
                rows = [r for r in reader if r]
                if not rows:
                    return {
                        "is_corrupted": True,
                        "status": "corrupted",
                        "reason": "CSV contains no parseable columns or rows.",
                        "details": {}
                    }

            return {
                "is_corrupted": False,
                "status": "intact",
                "reason": f"Character encoding and document structure are clean and intact ({'CSV' if category == 'csv' else 'Text'}).",
                "details": {}
            }
        except Exception as e:
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": f"Corrupted text stream: {str(e)}",
                "details": {"error": str(e)}
            }

    @classmethod
    def _check_generic_integrity(cls, file_path: str) -> dict:
        """Fallback check for binary or other supported types."""
        size = os.path.getsize(file_path)
        if size < 4:
            return {
                "is_corrupted": True,
                "status": "corrupted",
                "reason": "File is truncated (less than minimum 4-byte header size).",
                "details": {"size_bytes": size}
            }
        return {
            "is_corrupted": False,
            "status": "intact",
            "reason": "Binary data stream is readable and non-empty.",
            "details": {"size_bytes": size}
        }
