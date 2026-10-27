import os
import csv
from pathlib import Path
from pypdf import PdfReader
from pypdf.errors import PdfReadError
import docx
import openpyxl

class DocumentService:
    """
    Handles text extraction and content inspection across multiple document formats.
    Safely captures text for downstream AI classification, security scanning, and PII checks.
    """

    MAX_EXTRACT_CHARS = 50000  # Cap extracted text to prevent memory/context overflow

    @classmethod
    def extract_text(cls, file_path: str, detected_category: str) -> dict:
        """
        Extracts plain text and structure details from the file.
        Returns:
            {
                "success": bool,
                "text": str,
                "char_count": int,
                "word_count": int,
                "pages": int,
                "error": str | None,
                "is_encrypted": bool
            }
        """
        if not os.path.exists(file_path):
            return cls._result_error("File not found.")

        try:
            if detected_category == "pdf":
                return cls._extract_pdf(file_path)
            elif detected_category == "word":
                return cls._extract_docx(file_path)
            elif detected_category == "excel":
                return cls._extract_xlsx(file_path)
            elif detected_category in ["text", "csv"]:
                return cls._extract_plain_text(file_path)
            else:
                return {
                    "success": True,
                    "text": "",
                    "char_count": 0,
                    "word_count": 0,
                    "pages": 0,
                    "error": None,
                    "is_encrypted": False,
                    "note": f"Content extraction not supported for {detected_category} category."
                }
        except Exception as e:
            return cls._result_error(f"Text extraction failed: {str(e)}")

    @classmethod
    def _extract_pdf(cls, file_path: str) -> dict:
        # Pass 1: Try PyMuPDF (supports modern fonts, CID encodings, and layout blocks)
        try:
            import fitz
            doc = fitz.open(file_path)
            if doc.is_encrypted:
                doc.close()
                return {
                    "success": False,
                    "text": "",
                    "char_count": 0,
                    "word_count": 0,
                    "pages": 0,
                    "error": "Document is password-protected or encrypted.",
                    "is_encrypted": True
                }

            page_count = len(doc)
            mupdf_chunks = []
            total_chars = 0
            for page in doc:
                txt = (page.get_text("text") or "").strip()
                if txt:
                    mupdf_chunks.append(txt)
                    total_chars += len(txt)
                if total_chars >= cls.MAX_EXTRACT_CHARS:
                    break
            doc.close()

            full_text = "\n\n".join(mupdf_chunks)[:cls.MAX_EXTRACT_CHARS]
            
            # If PyMuPDF extracted substantial text, return it
            if len(full_text.strip()) >= 80 or (page_count <= 1 and full_text.strip()):
                return {
                    "success": True,
                    "text": full_text,
                    "char_count": len(full_text),
                    "word_count": len(full_text.split()),
                    "pages": page_count,
                    "error": None,
                    "is_encrypted": False
                }
        except Exception:
            full_text = ""
            page_count = 1

        # Pass 2: Fallback with PdfReader
        try:
            reader = PdfReader(file_path)
            if reader.is_encrypted:
                return {
                    "success": False,
                    "text": "",
                    "char_count": 0,
                    "word_count": 0,
                    "pages": len(reader.pages),
                    "error": "Document is password-protected or encrypted.",
                    "is_encrypted": True
                }

            extracted_chunks = []
            total_chars = 0
            for page in reader.pages:
                text = (page.extract_text() or "").strip()
                if text:
                    extracted_chunks.append(text)
                    total_chars += len(text)
                if total_chars >= cls.MAX_EXTRACT_CHARS:
                    break

            reader_text = "\n\n".join(extracted_chunks)[:cls.MAX_EXTRACT_CHARS]
            if len(reader_text.strip()) > len(full_text.strip()):
                full_text = reader_text
            page_count = max(page_count, len(reader.pages))
        except Exception:
            pass

        # Pass 3: Deep binary text salvage augmentation
        try:
            from services.pdf_recovery import PDFRecoveryEngine
            with open(file_path, "rb") as f:
                raw_bytes = f.read()
            salvaged = PDFRecoveryEngine.extract_all_text_from_pdf_binary(raw_bytes)
            if salvaged:
                seen_lines = set(line.strip().lower() for line in full_text.splitlines() if line.strip())
                new_lines = [s for s in salvaged if s.strip().lower() not in seen_lines]
                if new_lines:
                    if full_text.strip():
                        full_text = full_text.strip() + "\n\n" + "\n".join(new_lines)
                    else:
                        full_text = "\n".join(new_lines)
        except Exception:
            pass

        if full_text.strip():
            return {
                "success": True,
                "text": full_text[:cls.MAX_EXTRACT_CHARS],
                "char_count": len(full_text[:cls.MAX_EXTRACT_CHARS]),
                "word_count": len(full_text[:cls.MAX_EXTRACT_CHARS].split()),
                "pages": page_count if 'page_count' in locals() else 1,
                "error": None,
                "is_encrypted": False
            }

        return {
            "success": True,
            "text": "",
            "char_count": 0,
            "word_count": 0,
            "pages": page_count if 'page_count' in locals() else 1,
            "error": None,
            "is_encrypted": False
        }


    @classmethod
    def _extract_docx(cls, file_path: str) -> dict:
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n".join(paragraphs)[:cls.MAX_EXTRACT_CHARS]
            return {
                "success": True,
                "text": full_text,
                "char_count": len(full_text),
                "word_count": len(full_text.split()),
                "pages": max(1, len(full_text) // 2500),
                "error": None,
                "is_encrypted": False
            }
        except Exception as e:
            return cls._result_error(f"Unable to read Word document: {str(e)}")

    @classmethod
    def _extract_xlsx(cls, file_path: str) -> dict:
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            lines = []
            for sheet_name in wb.sheetnames[:5]:  # read up to first 5 sheets
                sheet = wb[sheet_name]
                lines.append(f"--- Sheet: {sheet_name} ---")
                row_count = 0
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
                    row_count += 1
                    if row_count > 500:
                        break
            wb.close()
            full_text = "\n".join(lines)[:cls.MAX_EXTRACT_CHARS]
            return {
                "success": True,
                "text": full_text,
                "char_count": len(full_text),
                "word_count": len(full_text.split()),
                "pages": len(wb.sheetnames),
                "error": None,
                "is_encrypted": False
            }
        except Exception as e:
            return cls._result_error(f"Unable to read Excel workbook: {str(e)}")

    @classmethod
    def _extract_plain_text(cls, file_path: str) -> dict:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read(cls.MAX_EXTRACT_CHARS)
            return {
                "success": True,
                "text": content,
                "char_count": len(content),
                "word_count": len(content.split()),
                "pages": max(1, len(content) // 3000),
                "error": None,
                "is_encrypted": False
            }
        except Exception as e:
            return cls._result_error(f"Unable to read text file: {str(e)}")

    @classmethod
    def _result_error(cls, message: str) -> dict:
        return {
            "success": False,
            "text": "",
            "char_count": 0,
            "word_count": 0,
            "pages": 0,
            "error": message,
            "is_encrypted": False
        }
