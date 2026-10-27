import os
import io
import pytest
import openpyxl
from services.xlsx_analyzer import XLSXAnalyzer
from services.xlsx_recovery import XLSXRecoveryEngine

def test_xlsx_header_displacement_recovery(tmp_path):
    # Create valid workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Units", "Amount"])
    ws.append(["Jan", 120, 1500.50])
    ws.append(["Feb", 95, 1200.00])

    buf = io.BytesIO()
    wb.save(buf)
    valid_bytes = buf.getvalue()

    # Corrupt with displaced header
    corrupt_bytes = b"CORRUPTED_PREFIX_BYTES_12345" + valid_bytes

    corrupt_path = str(tmp_path / "corrupt_sales.xlsx")
    out_path = str(tmp_path / "recovered_sales.xlsx")

    with open(corrupt_path, "wb") as f:
        f.write(corrupt_bytes)

    # Analyze
    analysis = XLSXAnalyzer.analyze(corrupt_path)
    assert analysis["is_corrupted"] is True
    assert any(r["region"] == "header" and r["status"] in ["warning", "corrupted"] for r in analysis["corruption_regions"])

    # Recover
    res = XLSXRecoveryEngine.recover(corrupt_path, out_path)
    assert res["success"] is True
    assert os.path.exists(out_path)

    # Verify with openpyxl
    wb_rec = openpyxl.load_workbook(out_path)
    assert "Sales" in wb_rec.sheetnames or len(wb_rec.sheetnames) >= 1

def test_xlsx_unrecoverable_empty(tmp_path):
    empty_path = str(tmp_path / "empty.xlsx")
    out_path = str(tmp_path / "out.xlsx")
    with open(empty_path, "wb") as f:
        pass

    res = XLSXRecoveryEngine.recover(empty_path, out_path)
    assert res["success"] is False
    assert res["status"] == "UNRECOVERABLE"
