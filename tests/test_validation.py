import os
import io
import tempfile
import pytest
from app import app
from services.ai_service import AIService
from services.validation_service import ValidationService

from config import Config

@pytest.fixture(autouse=True)
def isolate_test_history(tmp_path, monkeypatch):
    test_hist_file = str(tmp_path / "test_history.json")
    monkeypatch.setattr(Config, "HISTORY_FILE", test_hist_file)

@pytest.fixture
def client():
    app.config["TESTING"] = True
    with app.test_client() as client:
        yield client

def create_temp_file(content: bytes, suffix: str):
    """Helper to create a temporary test file."""
    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    tmp.write(content)
    tmp.close()
    return tmp.name

def create_valid_pdf_bytes() -> bytes:
    """Helper to create a valid, structurally intact minimal PDF."""
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()

# Test 1: Valid PDF below requested size -> PASS
def test_case_1_valid_pdf_below_size():
    pdf_content = create_valid_pdf_bytes()
    fpath = create_temp_file(pdf_content, ".pdf")
    
    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="contract.pdf",
            prompt="Validate this as a PDF document and make sure it is under 5 MB."
        )
        assert res["validation_status"] == "passed"
        assert res["score"] == 100
        assert res["file"]["is_corrupted"] is False
        assert res["file"]["integrity_status"] == "Intact"
        assert len(res["violations"]) == 0
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 2: PDF above requested size -> FAIL
def test_case_2_pdf_above_size():
    # Valid PDF bytes padded to exceed 50 KB
    from pypdf import PdfWriter
    writer = PdfWriter()
    for _ in range(50):
        writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    pdf_content = buf.getvalue()
    # Add dummy comment padding if needed
    if len(pdf_content) < 60 * 1024:
        pdf_content += b"\n%" + (b"P" * (60 * 1024))
    fpath = create_temp_file(pdf_content, ".pdf")

    try:
        # Prompt limits to 50 KB
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="large_doc.pdf",
            prompt="Make sure this PDF is below 50 KB."
        )
        assert res["validation_status"] == "failed"
        assert any("exceeds requested limit" in v for v in res["violations"])
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 3: EXE renamed to PDF -> FAIL (Disguised executable)
def test_case_3_exe_renamed_to_pdf():
    # Windows MZ executable header
    exe_content = b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00" + (b"\x00" * 200)
    fpath = create_temp_file(exe_content, ".pdf")

    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="innocent.pdf",
            prompt="Validate this PDF file under 10 MB."
        )
        assert res["validation_status"] == "failed"
        assert any("Dangerous" in v or "Executable" in v or "does not match" in v for v in res["violations"])
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 4: Correct extension but incorrect MIME / magic bytes -> FAIL
def test_case_4_extension_mismatch_content():
    # A plain text string saved as a .pdf
    fake_content = b"This is just plain text, not a real PDF binary!"
    fpath = create_temp_file(fake_content, ".pdf")

    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="fake.pdf",
            prompt="Validate this PDF document."
        )
        assert res["validation_status"] == "failed"
        assert any("does not match actual detected content" in v or "Expected PDF" in v for v in res["violations"])
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 5: Valid DOCX -> PASS
def test_case_5_valid_docx():
    import docx
    tmp_doc = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_doc.close()
    
    doc = docx.Document()
    doc.add_heading("Company Quarterly Review", level=1)
    doc.add_paragraph("This is a legitimate business report.")
    doc.save(tmp_doc.name)

    try:
        res = ValidationService.validate_upload(
            file_path=tmp_doc.name,
            original_filename="report.docx",
            prompt="Validate this Word document under 5 MB."
        )
        assert res["validation_status"] == "passed"
        assert res["score"] == 100
    finally:
        if os.path.exists(tmp_doc.name):
            os.remove(tmp_doc.name)

# Test 6: Unsupported file -> FAIL
def test_case_6_unsupported_file():
    # Arbitrary binary blob with unknown extension
    unknown_content = b"\x01\x02\x03\x04\x05\x06\x07\x08\x09\x10\x11\x12"
    fpath = create_temp_file(unknown_content, ".unknownbin")

    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="file.unknownbin",
            prompt="Check this file."
        )
        assert res["validation_status"] == "failed"
        assert any("Unsupported" in v or "Unverified" in v for v in res["violations"])
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 7: No attachment -> ERROR / 400
def test_case_7_no_attachment(client):
    response = client.post("/api/validate", data={"prompt": "Validate something"}, content_type="multipart/form-data")
    assert response.status_code == 400
    data = response.get_json()
    assert data["success"] is False
    assert "No file was attached" in data["violations"][0]

# Test 8: Prompt with no size requirement -> Perform only requested validations
def test_case_8_no_size_requirement():
    pdf_content = create_valid_pdf_bytes()
    fpath = create_temp_file(pdf_content, ".pdf")

    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="terms.pdf",
            prompt="Check whether this is a valid PDF."
        )
        assert res["validation_status"] == "passed"
        assert res["requirements"]["max_size_bytes"] is None
        assert res["file"]["is_corrupted"] is False
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 9: Prompt: "Accept PDF under 5 MB." -> PDF + 5 MB rule
def test_case_9_prompt_rule_extraction():
    parsed = AIService.parse_validation_prompt("Accept PDF under 5 MB.")
    assert "pdf" in parsed["allowed_types"]
    assert parsed["max_size_bytes"] == 5 * 1024 * 1024

# Test 10: Verify File Type, File Size, and File Integrity are checked in deterministic breakdown
def test_case_10_deterministic_checks():
    pdf_content = create_valid_pdf_bytes()
    fpath = create_temp_file(pdf_content, ".pdf")

    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="sample.pdf",
            prompt="Validate this as a PDF document and make sure it is under 5 MB."
        )
        assert res["validation_status"] == "passed"
        check_names = [c["name"] for c in res["checks"]]
        assert "File Type" in check_names
        assert "File Size" in check_names
        assert "File Integrity" in check_names
        assert len(res["checks"]) == 3
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 11: Corrupted PDF rejected -> FAIL
def test_case_11_corrupted_pdf_rejection():
    # Damaged PDF: header without cross-reference table or catalog
    damaged_pdf = b"%PDF-1.5\ncorrupted bytes and broken structure without xref\n%%EOF"
    fpath = create_temp_file(damaged_pdf, ".pdf")

    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="damaged.pdf",
            prompt="Validate this PDF and check if corrupted."
        )
        assert res["validation_status"] == "failed"
        assert res["file"]["is_corrupted"] is True
        assert res["file"]["integrity_status"] == "Corrupted"
        assert any("File corruption detected" in v for v in res["violations"])
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 12: Empty 0-byte file rejected as corrupted -> FAIL
def test_case_12_empty_file_corruption_rejection():
    fpath = create_temp_file(b"", ".pdf")

    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="empty.pdf",
            prompt="Validate this file."
        )
        assert res["validation_status"] == "failed"
        assert res["file"]["is_corrupted"] is True
        assert any("File is completely empty" in v or "File corruption" in v for v in res["violations"])
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

# Test 13: Corrupted DOCX (invalid ZIP structure) -> FAIL
def test_case_13_corrupted_docx_rejection():
    # Starts with PK header but has broken zip stream
    corrupt_zip = b"PK\x03\x04\x14\x00\x00\x00\x08\x00corrupted_payload"
    fpath = create_temp_file(corrupt_zip, ".docx")

    try:
        res = ValidationService.validate_upload(
            file_path=fpath,
            original_filename="corrupted.docx",
            prompt="Validate this Word document."
        )
        assert res["validation_status"] == "failed"
        assert res["file"]["is_corrupted"] is True
        assert any("corruption" in v.lower() for v in res["violations"])
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_delete_single_history_record(client):
    from services.history_service import HistoryService
    val_id = HistoryService.record_validation({
        "filename": "test_doc.pdf",
        "detected_type": "PDF",
        "file_size_human": "1.2 KB",
        "status": "passed",
        "score": 100
    })
    
    records = HistoryService.get_recent_validations()
    assert any(r.get("id") == val_id for r in records)

    # Test DELETE endpoint
    del_resp = client.delete(f"/api/history/{val_id}")
    assert del_resp.status_code == 200
    assert del_resp.get_json()["success"] is True

    # Verify removed
    records_after = HistoryService.get_recent_validations()
    assert not any(r.get("id") == val_id for r in records_after)

    # Test 404 on deleting non-existent ID
    del_resp_404 = client.delete("/api/history/non_existent_id")
    assert del_resp_404.status_code == 404

