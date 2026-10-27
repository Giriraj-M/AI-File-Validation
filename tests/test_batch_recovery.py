import os
import io
import pytest
from pathlib import Path
from app import app
from services.recovery_history_service import RecoveryHistoryService
import openpyxl

@pytest.fixture
def client():
    app.config["TESTING"] = True
    with app.test_client() as client:
        yield client

def test_multi_file_recovery_and_deletion(client):
    """
    Tests uploading 5 different files simultaneously to /api/recover/batch.
    Verifies independent results, history generation, and per-file record deletion.
    """
    # 1. Create a corrupt PDF (damaged xref/header)
    pdf_data = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>\nendobj\n4 0 obj\n<< /Length 44 >>\nstream\nBT /F1 12 Tf 100 700 Td (Batch PDF Test) Tj ET\nendstream\nendobj\n%%EOF"
    corrupt_pdf = b"GARBAGE_PREFIX" + pdf_data

    # 2. Create a corrupt XLSX
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestData"
    ws.cell(row=1, column=1, value="ID")
    ws.cell(row=1, column=2, value="Revenue")
    ws.cell(row=2, column=1, value=101)
    ws.cell(row=2, column=2, value=5000)
    xlsx_buf = io.BytesIO()
    wb.save(xlsx_buf)
    corrupt_xlsx = b"NET_GARBAGE" + xlsx_buf.getvalue()

    # 3. Create a damaged TXT
    corrupt_txt = b"Hello\x00World\x00Batch\x00Recovery"

    # 4. Create an unrecoverable 0-byte file
    unrec_file = b""

    # 5. Create a CSV file
    csv_file = b"Name,Age,Dept\nAlice,30,Engineering\nBob,28,Design"

    data = {
        "files[]": [
            (io.BytesIO(corrupt_pdf), "doc1.pdf"),
            (io.BytesIO(corrupt_xlsx), "data.xlsx"),
            (io.BytesIO(corrupt_txt), "notes.txt"),
            (io.BytesIO(unrec_file), "empty.dat"),
            (io.BytesIO(csv_file), "team.csv")
        ],
        "prompt": "Recover all files"
    }

    resp = client.post("/api/recover/batch", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200
    batch_result = resp.get_json()

    assert batch_result["success"] is True
    assert batch_result["total_files"] == 5
    assert len(batch_result["results"]) == 5

    # Check that individual files were accurately evaluated
    statuses = [r["status"] for r in batch_result["results"]]
    assert "FULLY_RECOVERED" in statuses or "SUCCESS" in statuses
    
    # The 0-byte file must NOT be reported as successfully recovered
    empty_result = next(r for r in batch_result["results"] if r["original_filename"] == "empty.dat")
    assert empty_result["status"] in ["UNRECOVERABLE", "FAILED"]
    assert empty_result["recovery_score"] == 0

    # Verify history stats calculation
    stats = RecoveryHistoryService.get_stats()
    assert stats["total_recovered"] > 0

    # Test individual file deletion
    target_rec = next(r for r in batch_result["results"] if r.get("recovery_id"))
    del_id = target_rec["recovery_id"]
    del_resp = client.post(f"/api/recovery/{del_id}/delete")
    assert del_resp.status_code == 200
    assert del_resp.get_json()["success"] is True

