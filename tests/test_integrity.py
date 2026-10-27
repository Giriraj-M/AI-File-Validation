import os
import io
import tempfile
import pytest
from pypdf import PdfWriter
import docx
import openpyxl
from PIL import Image
from services.integrity_service import IntegrityService
from services.ai_service import AIService

def create_temp_file(content: bytes, suffix: str) -> str:
    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    tmp.write(content)
    tmp.close()
    return tmp.name

def test_valid_pdf_integrity():
    w = PdfWriter()
    w.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    w.write(buf)
    
    fpath = create_temp_file(buf.getvalue(), ".pdf")
    try:
        res = IntegrityService.check_integrity(fpath, "pdf", "doc.pdf")
        assert res["is_corrupted"] is False
        assert res["status"] == "intact"
        assert "intact" in res["reason"].lower()
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_corrupted_pdf_integrity():
    damaged_bytes = b"%PDF-1.4\nthis is truncated and corrupted PDF data"
    fpath = create_temp_file(damaged_bytes, ".pdf")
    try:
        res = IntegrityService.check_integrity(fpath, "pdf", "damaged.pdf")
        assert res["is_corrupted"] is True
        assert res["status"] == "corrupted"
        assert "corrupted" in res["reason"].lower()
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_empty_zero_byte_file_integrity():
    fpath = create_temp_file(b"", ".pdf")
    try:
        res = IntegrityService.check_integrity(fpath, "pdf", "empty.pdf")
        assert res["is_corrupted"] is True
        assert res["status"] == "corrupted"
        assert "empty" in res["reason"].lower()      
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_valid_docx_integrity():
    fpath = create_temp_file(b"", ".docx")
    doc = docx.Document()
    doc.add_paragraph("Integrity test paragraph.")
    doc.save(fpath)
    try:
        res = IntegrityService.check_integrity(fpath, "word", "test.docx")
        assert res["is_corrupted"] is False
        assert res["status"] == "intact"
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_corrupted_docx_integrity():
    # Corrupted ZIP header
    fpath = create_temp_file(b"PK\x03\x04random_bad_data", ".docx")
    try:
        res = IntegrityService.check_integrity(fpath, "word", "corrupted.docx")
        assert res["is_corrupted"] is True
        assert res["status"] == "corrupted"
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_valid_xlsx_integrity():
    fpath = create_temp_file(b"", ".xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Revenue"
    ws["B1"] = 5000
    wb.save(fpath)
    wb.close()
    try:
        res = IntegrityService.check_integrity(fpath, "excel", "sheets.xlsx")
        assert res["is_corrupted"] is False
        assert res["status"] == "intact"
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_corrupted_xlsx_integrity():
    fpath = create_temp_file(b"PK\x03\x04truncated_xlsx_content", ".xlsx")
    try:
        res = IntegrityService.check_integrity(fpath, "excel", "bad.xlsx")
        assert res["is_corrupted"] is True
        assert res["status"] == "corrupted"
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_valid_image_integrity():
    img = Image.new("RGB", (50, 50), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    fpath = create_temp_file(buf.getvalue(), ".png")
    try:
        res = IntegrityService.check_integrity(fpath, "image", "sample.png")
        assert res["is_corrupted"] is False
        assert res["status"] == "intact"
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_corrupted_image_integrity():
    # Write PNG signature but truncate header/IHDR
    fpath = create_temp_file(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rtruncated", ".png")
    try:
        res = IntegrityService.check_integrity(fpath, "image", "broken.png")
        assert res["is_corrupted"] is True
        assert res["status"] == "corrupted"
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_null_byte_text_file_corruption():
    # Text file poisoned with null bytes
    fpath = create_temp_file(b"Valid text header\x00\x00corrupt binary", ".txt")
    try:
        res = IntegrityService.check_integrity(fpath, "text", "poisoned.txt")
        assert res["is_corrupted"] is True
        assert "null byte" in res["reason"].lower()
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)

def test_prompt_parsing_corruption_triggers():
    prompt = "Validate this PDF and check if the file is corrupted or damaged."
    parsed = AIService.parse_validation_prompt(prompt)
    assert parsed["check_integrity"] is True
    assert "pdf" in parsed["allowed_types"]

def test_unknown_binary_named_pdf_corrupted():
    # Arbitrary binary data without %PDF- magic bytes, named file.pdf
    fpath = create_temp_file(b"\x01\x02\x03\x04" * 50, ".pdf")
    try:
        res = IntegrityService.check_integrity(fpath, "unknown", "file.pdf", ["pdf"])
        assert res["is_corrupted"] is True
        assert res["status"] == "corrupted"
        assert "Missing valid '%PDF-' header signature" in res["reason"] or "corrupted" in res["reason"].lower()
    finally:
        if os.path.exists(fpath):
            os.remove(fpath)
